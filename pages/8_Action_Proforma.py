import streamlit as st
import json
import os
import sys
from datetime import datetime
import pytz
import io

# --- SAFE DOCX IMPORT ---
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_READY = True
except ImportError:
    DOCX_READY = False

# --- IMPORT GLOBAL CLOCK & LOGGER ---
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from utils import display_header_clocks, write_log

st.set_page_config(page_title="Action Proforma", page_icon="🚨", layout="wide")
# --- DATA MANAGEMENT ---
PROFORMA_FILE = "action_proforma_templates.json"

def load_proformas():
    """Loads templates from file or creates defaults if missing."""
    if os.path.exists(PROFORMA_FILE):
        with open(PROFORMA_FILE, "r") as f:
            return json.load(f)
    else:
        defaults = {
            "VIP MOVEMENT": ["Inform TM/APD", "Inform Security", "Brief Liaison Officer", "Verify Lounge", "Inform Fire Stn"],
            "ANTI-HIJACK": ["Activate Alarm", "Inform Committee", "Isolate Aircraft", "Coordinate Security", "Maintain Log"],
            "OVERDUE AIRCRAFT": ["Verify Neighbors", "Contact LFA", "Guard Freq (121.5)", "Inform SAR", "Inform Ops Officer"]
        }
        with open(PROFORMA_FILE, "w") as f:
            json.dump(defaults, f)
        return defaults

def save_proformas():
    """Saves the current session state templates to the JSON file."""
    with open(PROFORMA_FILE, "w") as f:
        json.dump(st.session_state.proformas, f)

def generate_docx(prof_name, items, state_data, controller):
    """Generates a formatted Word Document for the proforma."""
    doc = Document()
    title = doc.add_heading('ATC CHANDIGARH - ACTION PROFORMA', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph()
    p.add_run(f'PROTOCOL: {prof_name}').bold = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    info = doc.add_paragraph()
    info.add_run(f'Date: {datetime.now().strftime("%d %b %Y")}\n')
    info.add_run(f'Shift Controller: {controller}')
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Action Required'
    hdr_cells[1].text = 'Time (IST)'
    hdr_cells[2].text = 'Action Taken By'
    
    for i, itm in enumerate(items):
        entry = state_data.get(i, {"time": "PENDING", "name": "---"})
        row_cells = table.add_row().cells
        row_cells[0].text = itm
        row_cells[1].text = entry['time']
        row_cells[2].text = entry['name']
        
    doc.add_paragraph('\n\nVerified By: ____________________          Signature: ____________________')
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# Initialize Session State
if 'proformas' not in st.session_state:
    st.session_state.proformas = load_proformas()

# --- UI START ---
display_header_clocks()
st.title("🚨 Action Proforma Tracker")

raw_controller = st.session_state.get('duty_controllers', '').split(',')[0].strip().upper()
if not raw_controller:
    raw_controller = "UNKNOWN"

tab_execute, tab_manage = st.tabs(["🚀 Execute & Save", "⚙️ Manage Templates"])

# ==========================================
# TAB 1: EXECUTE & SAVE
# ==========================================
with tab_execute:
    if not st.session_state.proformas:
        st.info("No templates found. Please go to 'Manage Templates' to create one.")
    else:
        prof_names = sorted(list(st.session_state.proformas.keys()))
        selected_prof = st.selectbox("Select Action Proforma", prof_names)
        
        state_key = f"exe_state_{selected_prof}"
        if state_key not in st.session_state:
            st.session_state[state_key] = {}

        st.subheader(f"📋 Current Execution: {selected_prof}")
        items = st.session_state.proformas[selected_prof]
        
        h1, h2, h3, h4 = st.columns([2.5, 0.8, 1.2, 1.5])
        h1.markdown("**Action Required**")
        h2.markdown("**Done**")
        h3.markdown("**Time (IST)**")
        h4.markdown("**Action Taken By**")
        st.markdown("---")

        for idx, item in enumerate(items):
            c_txt, c_chk, c_time, c_name = st.columns([2.5, 0.8, 1.2, 1.5])
            c_txt.write(f"**{idx+1}.** {item}")
            is_checked = c_chk.checkbox("✓", key=f"chk_{selected_prof}_{idx}", label_visibility="collapsed")
            
            if is_checked:
                if idx not in st.session_state[state_key]:
                    ist_now = datetime.now(pytz.utc).astimezone(pytz.timezone('Asia/Kolkata')).strftime("%H:%M:%S")
                    st.session_state[state_key][idx] = {"time": ist_now, "name": raw_controller}
                    write_log("PROFORMA", f"{selected_prof}: {item} - COMPLETED BY {raw_controller} AT {ist_now}")

                current_entry = st.session_state[state_key][idx]
                c_time.success(f"🕒 {current_entry['time']}")
                new_name = c_name.text_input("Name", value=current_entry["name"], key=f"nm_{selected_prof}_{idx}", label_visibility="collapsed").upper()
                if new_name != current_entry["name"]:
                    st.session_state[state_key][idx]["name"] = new_name
            else:
                if idx in st.session_state[state_key]:
                    del st.session_state[state_key][idx]
                c_time.write("---")
                c_name.write("---")

        st.markdown("---")
        
        col_save, col_reset = st.columns(2)
        with col_save:
            docx_data = generate_docx(selected_prof, items, st.session_state[state_key], raw_controller)
            st.download_button(
                label="💾 Save as Word Document (.docx)",
                data=docx_data,
                file_name=f"Proforma_{selected_prof}_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                type="primary"
            )
        with col_reset:
            if st.button("🏁 Reset Proforma Screen", use_container_width=True):
                st.session_state[state_key] = {}
                st.rerun()

# ==========================================
# TAB 2: MANAGE TEMPLATES
# ==========================================
with tab_manage:
    st.subheader("⚙️ Proforma Template Management")
    
    col_m1, col_m2 = st.columns([1, 1.5])
    
    with col_m1:
        st.markdown("**➕ Create New Template**")
        with st.container(border=True):
            new_p_name = st.text_input("Enter Proforma Name (e.g., BOMB THREAT)")
            if st.button("Create Proforma", type="primary", use_container_width=True):
                if new_p_name:
                    clean_name = new_p_name.strip().upper()
                    if clean_name not in st.session_state.proformas:
                        st.session_state.proformas[clean_name] = []
                        save_proformas()
                        st.success(f"Created: {clean_name}")
                        st.rerun()
                    else:
                        st.warning("Template already exists.")

    with col_m2:
        st.markdown("**✏️ Edit/Delete Steps**")
        if not st.session_state.proformas:
            st.info("No templates available to edit.")
        else:
            with st.container(border=True):
                all_keys = sorted(list(st.session_state.proformas.keys()))
                edit_target = st.selectbox("Select Template to Edit", all_keys)
                
                # Add New Step
                st.write(f"---")
                new_step = st.text_input(f"Add new step to {edit_target}:", placeholder="e.g., Inform Duty Officer")
                if st.button("Add Step", use_container_width=True):
                    if new_step:
                        st.session_state.proformas[edit_target].append(new_step.strip())
                        save_proformas()
                        st.rerun()
                
                # List Existing Steps
                st.write(f"**Current Steps in {edit_target}:**")
                steps = st.session_state.proformas[edit_target]
                if not steps:
                    st.write("*No steps added yet.*")
                else:
                    for i, step in enumerate(steps):
                        sc1, sc2 = st.columns([5, 1])
                        sc1.write(f"{i+1}. {step}")
                        if sc2.button("🗑️", key=f"del_{edit_target}_{i}"):
                            st.session_state.proformas[edit_target].pop(i)
                            save_proformas()
                            st.rerun()
                
                st.write("---")
                # Delete Entire Template
                with st.expander("🚨 Danger Zone: Delete Entire Template"):
                    if st.button(f"Permanently Delete '{edit_target}'", type="primary", use_container_width=True):
                        del st.session_state.proformas[edit_target]
                        save_proformas()
                        st.rerun()